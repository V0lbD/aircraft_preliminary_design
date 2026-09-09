from __future__ import annotations

import pytest
import openpyxl

from aircraft_design.io.excel_loader import load_standard_table, load_global_coefficients
from aircraft_design.core.errors import FileFormatError


def test_load_standard_table_reads_correct_cells(tmp_path) -> None:
    # 1. Подготавливаем фейковый Excel-файл
    wb = openpyxl.Workbook()
    ws = wb.active

    # Имитируем крыло, первая алюминиевая технология (строка 3, столбцы B, C, D)
    ws.cell(row=3, column=2, value=2.0)  # Обшивка
    ws.cell(row=3, column=3, value=5.0)  # Поперечный
    ws.cell(row=3, column=4, value=1.0)  # Продольный

    # Имитируем оперение, вторая композитная технология (строка 6, столбцы H, I, J)
    ws.cell(row=6, column=8, value=1.5)
    ws.cell(row=6, column=9, value=3.0)
    ws.cell(row=6, column=10, value=0.5)

    file_path = tmp_path / "table1_mock.xlsx"
    wb.save(file_path)

    # 2. Вызываем наш загрузчик
    table_data = load_standard_table(file_path)

    # 3. Проверяем результаты
    assert table_data.wing.alum_1.skin == 2.0
    assert table_data.wing.alum_1.longitudinal == 1.0  # D3
    assert table_data.wing.alum_1.transverse == 5.0  # C3

    assert table_data.tail.comp_2.skin == 1.5
    assert table_data.tail.comp_2.transverse == 3.0

    # Не заполненные ячейки должны стать 0.0
    assert table_data.fuselage.alum_1.skin == 0.0


def test_load_global_coefficients_reads_correct_cells(tmp_path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active

    # Имитируем значения в столбце А (колонка 1)
    values = [0.10, 0.07, 0.10, 0.12, 0.15, 1.0, 0.20, 4.0, 110000.0, 1.20]
    for i, val in enumerate(values, start=1):
        ws.cell(row=i, column=1, value=val)

    # Добавим строковое значение с запятой, чтобы проверить наш парсер (частая беда Excel)
    ws.cell(row=9, column=1, value="110000,50")

    file_path = tmp_path / "table10_mock.xlsx"
    wb.save(file_path)

    # 2. Вызываем загрузчик
    globals_data = load_global_coefficients(file_path)

    # 3. Проверяем результаты
    assert globals_data.k_skl == 0.10
    assert globals_data.k_oper == 0.07
    assert globals_data.c_1m == 110000.50  # Запятая должна была корректно превратиться в точку
    assert globals_data.k_prib == 1.20


def test_loader_raises_error_on_missing_file() -> None:
    with pytest.raises(FileFormatError):
        load_standard_table("non_existent_file.xlsx")