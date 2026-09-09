from __future__ import annotations

from pathlib import Path
import openpyxl
import re

from aircraft_design.core.models.technology import (
    ComponentTechnologies,
    TableData,
    TechCoefficients,
    GlobalCoefficients,
    TechnologyDatabase,
)
from aircraft_design.core.errors import FileFormatError


def _parse_component(sheet: openpyxl.worksheet.worksheet.Worksheet, start_row: int,
                     start_col: int) -> ComponentTechnologies:
    """
    Читает блок 4x3 ячейки (4 технологии, 3 элемента конструкции).
    start_col: B=2, E=5, H=8
    """

    def _read_tech(r: int) -> TechCoefficients:
        # Извлекаем значения, преобразуя запятые в точки (на случай кривого формата в Excel)
        val_skin = str(sheet.cell(row=r, column=start_col).value or 0.0).replace(',', '.')
        val_trans = str(sheet.cell(row=r, column=start_col + 1).value or 0.0).replace(',', '.')
        val_long = str(sheet.cell(row=r, column=start_col + 2).value or 0.0).replace(',', '.')

        return TechCoefficients(
            skin=float(val_skin),
            longitudinal=float(val_long),
            transverse=float(val_trans)
        )

    return ComponentTechnologies(
        alum_1=_read_tech(start_row),
        alum_2=_read_tech(start_row + 1),
        comp_1=_read_tech(start_row + 2),
        comp_2=_read_tech(start_row + 3),
    )


def load_standard_table(filepath: str | Path) -> TableData:
    """Загружает данные из стандартной таблицы (№ 1-9)."""
    path = Path(filepath)
    if not path.is_file():
        raise FileFormatError(f"Файл таблицы не найден: {path}")

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = wb.active

        return TableData(
            wing=_parse_component(sheet, start_row=3, start_col=2),  # Столбцы B, C, D
            fuselage=_parse_component(sheet, start_row=3, start_col=5),  # Столбцы E, F, G
            tail=_parse_component(sheet, start_row=3, start_col=8),  # Столбцы H, I, J
        )
    finally:
        if 'wb' in locals():
            wb.close()


def load_global_coefficients(filepath: str | Path) -> GlobalCoefficients:
    """Загружает Таблицу 10."""
    path = Path(filepath)
    if not path.is_file():
        raise FileFormatError(f"Файл таблицы не найден: {path}")

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        sheet = wb.active

        def _val(row: int) -> float:
            # Значения находятся в столбце А (колонка 1)
            raw_val = sheet.cell(row=row, column=1).value
            if raw_val is None:
                return 0.0
            return float(str(raw_val).replace(',', '.'))

        return GlobalCoefficients(
            k_skl=_val(1),
            k_oper=_val(2),
            k_prod=_val(3),
            k_ras=_val(4),
            k_vspom=_val(5),
            k_sb=_val(6),
            k_adm=_val(7),
            s_k_adm=_val(8),
            c_1m=_val(9),
            k_prib=_val(10),
        )
    except Exception as e:
        raise FileFormatError(f"Ошибка чтения файла {path.name}: {e}")
    finally:
        if 'wb' in locals():
            wb.close()


def load_technology_database(tables_dir: str | Path) -> TechnologyDatabase:
    """
    Сканирует папку с таблицами, находит 10 нужных файлов по цифре в начале имени
    и собирает их в единую базу данных TechnologyDatabase.
    """
    dir_path = Path(tables_dir)
    if not dir_path.is_dir():
        raise FileFormatError(f"Директория с таблицами не найдена: {dir_path}")

    file_map = {}
    for f in dir_path.glob("*.xlsx"):
        # Ищем числа в начале имени файла (например, "1", "10")
        match = re.match(r"^(\d+)", f.name)
        if match:
            file_map[match.group(1)] = f

    def get_file(num: str) -> Path:
        if num not in file_map:
            raise FileFormatError(f"Не найден Excel-файл для Таблицы {num} в папке {dir_path}")
        return file_map[num]

    return TechnologyDatabase(
        nel=load_standard_table(get_file("1")),
        ndet=load_standard_table(get_file("2")),
        m_el=load_standard_table(get_file("3")),
        kim=load_standard_table(get_file("4")),
        c1=load_standard_table(get_file("5")),
        cst1=load_standard_table(get_file("6")),
        nrb=load_standard_table(get_file("7")),
        cnrb=load_standard_table(get_file("8")),
        s1=load_standard_table(get_file("9")),
        globals=load_global_coefficients(get_file("10")),
    )